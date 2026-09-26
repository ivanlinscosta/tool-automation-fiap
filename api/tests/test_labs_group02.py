import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.config import settings


LAB_GROUP = "test-labs-02"
INSTRUCTOR_KEY = "chave-de-instrutor"
BASE = "/api/v1/labs/groups"


def _headers(**extra: str) -> dict[str, str]:
    return {"X-Lab-Group": LAB_GROUP, **extra}


@pytest.fixture(autouse=True)
def _instructor_key(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "LABS_INSTRUCTOR_KEY", INSTRUCTOR_KEY)


def _list_rates(client: TestClient, group: str = "02", **params: object) -> dict:
    response = client.get(f"{BASE}/{group}/rates", params=params, headers=_headers())
    assert response.status_code == 200, response.text
    return response.json()


def _list_files(client: TestClient, group: str = "02", **params: object) -> dict:
    response = client.get(f"{BASE}/{group}/files", params=params, headers=_headers())
    assert response.status_code == 200, response.text
    return response.json()


def test_group02_seeds_at_least_one_thousand_rates(client: TestClient) -> None:
    payload = _list_rates(client, limit=1)
    assert payload["meta"]["total"] >= 1000


def test_group02_lazy_seed_happens_once(client: TestClient) -> None:
    first = _list_rates(client, limit=1)
    second = _list_rates(client, limit=1)
    assert first["meta"]["total"] == second["meta"]["total"]


def test_group02_pagination_is_stable_and_non_overlapping(client: TestClient) -> None:
    first = _list_rates(client, limit=10, offset=0, sort="rate_id")
    second = _list_rates(client, limit=10, offset=10, sort="rate_id")
    assert len(first["items"]) == 10
    assert len(second["items"]) == 10
    assert {item["rate_id"] for item in first["items"]}.isdisjoint(
        {item["rate_id"] for item in second["items"]}
    )


def test_group02_offset_beyond_total_returns_empty_page(client: TestClient) -> None:
    payload = _list_rates(client, limit=5, offset=10_000_000)
    assert payload["items"] == []
    assert payload["meta"]["has_more"] is False


@pytest.mark.parametrize("alias", ["2", "02", "group-2", "group-02", "GROUP-02"])
def test_group02_accepts_every_group_alias(client: TestClient, alias: str) -> None:
    assert _list_rates(client, group=alias, limit=1)["meta"]["total"] >= 1000


@pytest.mark.parametrize("bad", ["0", "13", "99", "abc", "group-13", "group-0"])
def test_group02_rejects_unknown_group_alias(client: TestClient, bad: str) -> None:
    response = client.get(f"{BASE}/{bad}/rates", headers=_headers())
    assert response.status_code == 404
    assert "Unknown lab group" in response.json()["detail"]


def test_group02_endpoint_rejects_other_group(client: TestClient) -> None:
    response = client.get(f"{BASE}/03/rates", headers=_headers())
    assert response.status_code == 404
    assert "does not belong to lab group 02" in response.json()["detail"]


def test_group02_filters_by_status_indexador_mes_and_search(client: TestClient) -> None:
    invalid_rows = _list_rates(client, limit=50, status="INVALIDO")
    assert invalid_rows["meta"]["total"] > 0
    assert all(item["status"] == "INVALIDO" for item in invalid_rows["items"])

    sample = invalid_rows["items"][0]
    by_indexador = _list_rates(client, limit=50, indexador=sample["indexador"])
    assert all(
        item["indexador"] == sample["indexador"] for item in by_indexador["items"]
    )

    by_month = _list_rates(client, limit=50, mes=sample["mes"])
    assert all(item["mes"] == sample["mes"] for item in by_month["items"])

    searched = _list_rates(client, limit=25, search="cnpj")
    assert searched["meta"]["total"] > 0


def test_group02_sorting_is_deterministic_and_rejects_unknown_field(
    client: TestClient,
) -> None:
    ascending = _list_rates(client, limit=5, sort="taxa", order="asc")
    descending = _list_rates(client, limit=5, sort="taxa", order="desc")
    assert [item["taxa"] for item in ascending["items"]] == sorted(
        item["taxa"] for item in ascending["items"]
    )
    assert ascending["items"][0]["taxa"] != descending["items"][0]["taxa"]

    invalid = client.get(
        f"{BASE}/02/rates", params={"sort": "campo_inexistente"}, headers=_headers()
    )
    assert invalid.status_code == 400
    assert "Invalid sort field" in invalid.json()["detail"]


def test_group02_get_rate_by_id_and_missing_rate(client: TestClient) -> None:
    first = _list_rates(client, limit=1)["items"][0]
    ok = client.get(f"{BASE}/02/rates/{first['rate_id']}", headers=_headers())
    assert ok.status_code == 200
    assert ok.json()["rate_id"] == first["rate_id"]

    missing = client.get(f"{BASE}/02/rates/RATE-999999", headers=_headers())
    assert missing.status_code == 404
    assert missing.json()["detail"] == "PrimaryMarketRate 'RATE-999999' not found"


def test_group02_files_are_listed_and_can_be_created(client: TestClient) -> None:
    listed = _list_files(client, limit=5)
    assert listed["meta"]["total"] >= 20

    created = client.post(
        f"{BASE}/02/files",
        json={
            "nome_arquivo": "nova_carga_setembro.xlsx",
            "origem": "tesouraria",
            "status": "RECEIVED",
        },
        headers=_headers(),
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["status"] == "RECEIVED"
    assert body["records_valid"] == 0
    assert body["records_invalid"] == 0

    stored = client.get(f"{BASE}/02/files/{body['file_id']}", headers=_headers())
    assert stored.status_code == 200
    assert stored.json()["nome_arquivo"] == "nova_carga_setembro.xlsx"


def test_group02_file_download_is_xlsx_and_deterministic(client: TestClient) -> None:
    file_id = _list_files(client, limit=1, status="REJECTED")["items"][0]["file_id"]
    first = client.get(f"{BASE}/02/files/{file_id}/download", headers=_headers())
    second = client.get(f"{BASE}/02/files/{file_id}/download", headers=_headers())
    assert first.status_code == 200 and second.status_code == 200
    assert (
        first.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment; filename=" in first.headers["content-disposition"]
    assert first.content == second.content
    assert first.headers["content-length"] == str(len(first.content))

    workbook = load_workbook(io.BytesIO(first.content))
    assert workbook.active.max_row >= 2


def test_group02_rejected_files_have_consistent_counts(client: TestClient) -> None:
    files = _list_files(client, limit=200)["items"]
    rates = _list_rates(client, limit=200)["meta"]["total"]
    assert rates >= 1000
    counts: dict[str, int] = {}
    offset = 0
    total = _list_rates(client, limit=1)["meta"]["total"]
    while offset < total:
        for row in _list_rates(client, limit=200, offset=offset)["items"]:
            counts[row["file_id"]] = counts.get(row["file_id"], 0) + 1
        offset += 200
    rejected = [item for item in files if item["status"] == "REJECTED"]
    assert rejected
    for item in files:
        assert item["records_valid"] + item["records_invalid"] == counts.get(
            item["file_id"], 0
        )
    assert any(item["records_invalid"] > 0 for item in rejected)


def test_group02_temporal_coherence_across_rates_and_files(client: TestClient) -> None:
    total = _list_rates(client, limit=1)["meta"]["total"]
    offset = 0
    while offset < total:
        for item in _list_rates(client, limit=200, offset=offset)["items"]:
            assert item["created_at"] <= item["updated_at"]
        offset += 200
    for item in _list_files(client, limit=200)["items"]:
        assert item["created_at"] <= item["updated_at"]


def test_group02_integrity_endpoint_reports_healthy(client: TestClient) -> None:
    response = client.get(
        f"{BASE}/02/integrity", headers=_headers(**{"X-Instructor-Key": INSTRUCTOR_KEY})
    )
    assert response.status_code == 200
    body = response.json()
    assert body["healthy"] is True
    assert body["violations"] == []


@pytest.mark.parametrize(
    ("scenario", "expected"),
    [
        ("validation_error", 422),
        ("not_found", 404),
        ("duplicate", 409),
        ("server_error", 500),
    ],
)
def test_group02_scenarios_on_read(
    client: TestClient, scenario: str, expected: int
) -> None:
    response = client.get(
        f"{BASE}/02/rates", params={"scenario": scenario}, headers=_headers()
    )
    assert response.status_code == expected


@pytest.mark.parametrize(
    ("scenario", "expected"),
    [
        ("validation_error", 422),
        ("not_found", 404),
        ("duplicate", 409),
        ("server_error", 500),
    ],
)
def test_group02_scenarios_on_write(
    client: TestClient, scenario: str, expected: int
) -> None:
    response = client.post(
        f"{BASE}/02/files",
        params={"scenario": scenario},
        json={
            "nome_arquivo": "arquivo_teste.xlsx",
            "origem": "mesa_credito",
            "status": "RECEIVED",
        },
        headers=_headers(),
    )
    assert response.status_code == expected


def test_group02_timeout_unknown_and_success_scenarios(client: TestClient) -> None:
    timeout = client.get(
        f"{BASE}/02/rates", params={"scenario": "timeout"}, headers=_headers()
    )
    assert timeout.status_code == 504

    unknown = client.get(
        f"{BASE}/02/rates", params={"scenario": "explodir"}, headers=_headers()
    )
    assert unknown.status_code == 422

    success = _list_rates(client, limit=1, scenario="success")
    assert success["meta"]["total"] >= 1000


def test_group02_scenario_does_not_mutate_state(client: TestClient) -> None:
    before = _list_files(client, limit=1)["meta"]["total"]
    client.post(
        f"{BASE}/02/files",
        params={"scenario": "server_error"},
        json={
            "nome_arquivo": "nao_deve_criar.xlsx",
            "origem": "tesouraria",
            "status": "RECEIVED",
        },
        headers=_headers(),
    )
    after = _list_files(client, limit=1)["meta"]["total"]
    assert after == before


def test_group02_data_is_deterministic_across_reseeds(
    client: TestClient, db_session: Session
) -> None:
    from app.labs.registry import purge_group

    first = _list_rates(client, limit=200)["items"]
    purge_group(db_session, 2)
    second = _list_rates(client, limit=200)["items"]
    assert first == second
